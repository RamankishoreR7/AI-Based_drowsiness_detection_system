import cv2
import time
import numpy as np
import os
import platform
import winsound
import mediapipe as mp
import pywhatkit as kit
import threading

from mediapipe.tasks.python import vision
from mediapipe.tasks.python.core import base_options
from openpyxl import Workbook, load_workbook
from datetime import datetime
# 5-Second Emergency Alert Sound
def play_alert_sound():
    if IS_WINDOWS:
        end_time = time.time() + 5
        while time.time() < end_time:
            winsound.Beep(1500, 300)
            winsound.Beep(1000, 300)
# WhatsApp Alert
def send_whatsapp_alert():
    try:
        phone_number = "+919952215095"
        message = "⚠ ALERT: Driver drowsiness detected! Please check immediately."
        kit.sendwhatmsg_instantly(phone_number, message, wait_time=15, tab_close=True)
        print("WhatsApp alert sent")
    except Exception as e:
        print("WhatsApp alert failed:", e)
# Event Photo Capture
def save_event_photo(frame):
    folder = "drowsy_events"
    if not os.path.exists(folder):
        os.makedirs(folder)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{folder}/event_{timestamp}.jpg"
    cv2.imwrite(filename, frame)
    print(f"Event photo saved: {filename}")
# Excel Setup
EXCEL_FILE = "drowsiness_log.xlsx"

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Drowsiness Log"
        ws.append([
            "Event No", "Start Time", "End Time",
            "Duration (s)", "Session Time (s)", "Total Events"
        ])
        wb.save(EXCEL_FILE)

def log_event_to_excel(event_no, start, end, duration, session_time, total_events):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    ws.append([
        event_no, start, end,
        round(duration, 2),
        round(session_time, 2),
        total_events
    ])
    wb.save(EXCEL_FILE)

init_excel()
# Model path
MODEL_PATH = "face_landmarker.task"
if not os.path.exists(MODEL_PATH):
    raise FileNotFoundError("face_landmarker.task not found in project folder")
# Load FaceLandmarker
options = vision.FaceLandmarkerOptions(
    base_options=base_options.BaseOptions(model_asset_path=MODEL_PATH),
    running_mode=vision.RunningMode.VIDEO,
    num_faces=1
)
face_landmarker = vision.FaceLandmarker.create_from_options(options)
# Eye landmarks
LEFT_EYE = [33, 160, 158, 133, 153, 144]
RIGHT_EYE = [362, 385, 387, 263, 373, 380]

def eye_aspect_ratio(eye):
    A = np.linalg.norm(eye[1] - eye[5])
    B = np.linalg.norm(eye[2] - eye[4])
    C = np.linalg.norm(eye[0] - eye[3])
    if C == 0:
        return 0.0
    return (A + B) / (2.0 * C)
# Parameters

EAR_THRESHOLD = 0.15
EYE_CLOSED_TIME = 1.5
BEEP_INTERVAL = 0.3

# State variables
eye_closed_start = None
current_event_active = False
current_event_start = None
event_start_clock = None

drowsy_event_count = 0
total_drowsy_time = 0.0
live_eye_closed_time = 0.0

last_beep_time = 0
session_start_time = time.time()
IS_WINDOWS = platform.system() == "Windows"
# Camera
cap = cv2.VideoCapture(0)
cap.set(cv2.CAP_PROP_BUFFERSIZE, 1)
# Main loop

while True:
    ret, frame = cap.read()
    if not ret:
        break

    frame = cv2.flip(frame, 1)
    rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
    timestamp_ms = int(time.time() * 1000)

    mp_image = mp.Image(image_format=mp.ImageFormat.SRGB, data=rgb)
    result = face_landmarker.detect_for_video(mp_image, timestamp_ms)

    if result.face_landmarks:
        landmarks = result.face_landmarks[0]
        h, w, _ = frame.shape

        left_eye = np.array([[int(landmarks[i].x * w), int(landmarks[i].y * h)] for i in LEFT_EYE])
        right_eye = np.array([[int(landmarks[i].x * w), int(landmarks[i].y * h)] for i in RIGHT_EYE])

        ear = (eye_aspect_ratio(left_eye) + eye_aspect_ratio(right_eye)) / 2.0

        cv2.putText(frame, f"EAR: {ear:.2f}", (30, 40),
                    cv2.FONT_HERSHEY_SIMPLEX, 1, (0, 255, 0), 2)

        if ear < EAR_THRESHOLD:
            if eye_closed_start is None:
                eye_closed_start = time.time()

            live_eye_closed_time = time.time() - eye_closed_start

            if live_eye_closed_time >= EYE_CLOSED_TIME:
                cv2.putText(frame, "DROWSINESS ALERT!", (40, 90),
                            cv2.FONT_HERSHEY_SIMPLEX, 1.1, (0, 0, 255), 3)

                if not current_event_active:
                    current_event_active = True
                    current_event_start = time.time()
                    event_start_clock = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    drowsy_event_count += 1

                    save_event_photo(frame)
                    threading.Thread(target=play_alert_sound, daemon=True).start()
                    threading.Thread(target=send_whatsapp_alert, daemon=True).start()

                now = time.time()
                if now - last_beep_time >= BEEP_INTERVAL:
                    if IS_WINDOWS:
                        winsound.Beep(1000, 200)
                    last_beep_time = now
        else:
            live_eye_closed_time = 0.0
            if current_event_active:
                event_end_time = time.time()
                duration = event_end_time - current_event_start
                total_drowsy_time += duration

                event_end_clock = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                session_time = time.time() - session_start_time

                log_event_to_excel(
                    drowsy_event_count,
                    event_start_clock,
                    event_end_clock,
                    duration,
                    session_time,
                    drowsy_event_count
                )

                current_event_active = False
                current_event_start = None
            eye_closed_start = None
    else:
        live_eye_closed_time = 0.0
        eye_closed_start = None

    session_time = time.time() - session_start_time
    status_text = "DROWSY" if current_event_active else "NORMAL"
    status_color = (0, 0, 255) if current_event_active else (0, 255, 0)

    h, w, _ = frame.shape
    px1, py1 = w - 420, h - 170
    px2, py2 = w - 10,  h - 10

    cv2.rectangle(frame, (px1, py1), (px2, py2), (0, 0, 0), -1)
    cv2.rectangle(frame, (px1, py1), (px2, py2), (255, 255, 255), 1)

    cv2.putText(frame, f"Status : {status_text}", (px1 + 10, py1 + 30),
                cv2.FONT_HERSHEY_SIMPLEX, 0.7, status_color, 2)

    cv2.imshow("Driver Drowsiness Detection (Real-Time)", frame)

    if cv2.waitKey(1) & 0xFF == 27:
        break

print("\n--- SESSION SUMMARY ---")
print(f"Drowsy events     : {drowsy_event_count}")
print(f"Total drowsy time : {total_drowsy_time:.2f} seconds")

cap.release()
cv2.destroyAllWindows()
face_landmarker.close()